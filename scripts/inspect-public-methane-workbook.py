"""Read-only XLSX extraction into a separate source-cell-addressed JSON file."""
import argparse
import datetime
import hashlib
import json
from pathlib import Path
import openpyxl

def encode(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return {'type': type(value).__name__, 'iso': value.isoformat()}
    return value

def main():
    p=argparse.ArgumentParser()
    p.add_argument('--input', type=Path, required=True)
    p.add_argument('--sha256', required=True)
    p.add_argument('--output', type=Path, required=True)
    a=p.parse_args()
    before=hashlib.sha256(a.input.read_bytes()).hexdigest()
    if before != a.sha256:
        raise ValueError('Input does not match acquisition receipt')
    book=openpyxl.load_workbook(a.input, read_only=True, data_only=False)
    cached=openpyxl.load_workbook(a.input, read_only=True, data_only=True)
    sheets=[]
    for sheet in book:
        values=cached[sheet.title]
        rows=[]
        for row_number, row in enumerate(sheet.iter_rows(), start=1):
            cells=[]
            for cell in row:
                if cell.value is not None:
                    record={'cell':cell.coordinate,'value':encode(cell.value),'dataType':cell.data_type,'numberFormat':cell.number_format}
                    if cell.data_type == 'f':
                        record['cachedValue']=encode(values[cell.coordinate].value)
                    cells.append(record)
            if cells:
                rows.append({'row':row_number,'cells':cells})
        sheets.append({'name':sheet.title,'maxRow':sheet.max_row,'maxColumn':sheet.max_column,'populatedRows':len(rows),'nonemptyCells':sum(len(r['cells']) for r in rows),'formulaCells':sum(c['dataType']=='f' for r in rows for c in r['cells']),'rows':rows})
    book.close();cached.close()
    if hashlib.sha256(a.input.read_bytes()).hexdigest()!=before:
        raise ValueError('Source changed during inspection')
    result={'schemaVersion':'public-methane-workbook-structure-v1','createdAt':datetime.datetime.now(datetime.timezone.utc).isoformat(),'source':str(a.input),'sha256':before,'sourceUnchanged':True,'openpyxlVersion':openpyxl.__version__,'scriptSha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'sheets':sheets,'recalculated':False,'scientificAdmission':False}
    a.output.parent.mkdir(parents=True, exist_ok=True)
    with a.output.open('x') as f:
        json.dump(result,f,indent=2,allow_nan=False)
        f.write('\n')
    print(json.dumps({**{k:v for k,v in result.items() if k!='sheets'},'sheets':[{**{k:v for k,v in s.items() if k!='rows'},'firstRows':s['rows'][:8]} for s in sheets]},indent=2))

if __name__=='__main__':
    main()
